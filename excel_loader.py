from pathlib import Path
from typing import Any, Iterator


HEADER_MAP = {
    "id": "id",
    "nguon": "source",
    "de bai": "title",
    "noi dung": "content",
    "boi canh": "context",
    "the loai bai": "category",
    "ky thuat su dung": "technique",
    "database_schema": "database_schema",
    "tu khoa": "keywords",
}


def _normalize_header(value: Any) -> str:
    import unicodedata

    text = unicodedata.normalize("NFD", str(value or "").strip().lower())
    return " ".join(
        "".join(char for char in text if unicodedata.category(char) != "Mn").split()
    )


def load_excel_rows(path: str, limit: int | None = None) -> Iterator[dict[str, Any]]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError(
            "openpyxl is required to read Excel files. "
            "Run: python -m pip install -r requirements.txt"
        ) from exc

    resolved = Path(path).resolve()
    if not resolved.exists():
        raise FileNotFoundError(f"Excel file not found: {resolved}")

    workbook = load_workbook(resolved, read_only=True, data_only=True)
    try:
        sheet = workbook.active
        rows = sheet.iter_rows(values_only=True)
        raw_headers = next(rows, None)
        if not raw_headers:
            raise ValueError("Excel sheet is empty")

        headers = [HEADER_MAP.get(_normalize_header(value)) for value in raw_headers]
        if "database_schema" not in headers:
            raise ValueError("Excel must contain a 'database_schema' column")

        emitted = 0
        for row_number, values in enumerate(rows, start=2):
            record = {
                header: value.strip() if isinstance(value, str) else value
                for header, value in zip(headers, values)
                if header is not None
            }
            if not any(value not in (None, "") for value in record.values()):
                continue
            record["excel_row"] = row_number
            yield record
            emitted += 1
            if limit is not None and emitted >= limit:
                break
    finally:
        workbook.close()


def get_excel_row(path: str, excel_id: str) -> dict[str, Any]:
    wanted = excel_id.strip().casefold()
    if not wanted:
        raise ValueError("excel_id must not be empty")
    for row in load_excel_rows(path):
        if str(row.get("id") or "").strip().casefold() == wanted:
            return row
    raise ValueError(f"Excel ID not found: {excel_id}")
